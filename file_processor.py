import os
import base64
import mimetypes
import uuid
import csv
from io import StringIO
from PIL import Image
import pypdf
import docx

UPLOAD_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

ALLOWED_EXTENSIONS = {
    "image": {"png", "jpg", "jpeg", "webp", "gif", "bmp", "svg", "tiff", "ico"},
    "audio": {"mp3", "wav", "m4a", "ogg", "webm", "flac", "aac", "wma"},
    "video": {"mp4", "webm", "mov", "mkv", "avi", "ogv", "wmv"},
    "document": {"pdf", "docx", "doc", "txt", "md", "csv", "tsv", "xlsx", "xls", "json", "xml", "html", "htm", "log", "rtf"},
    "code": {"py", "js", "ts", "jsx", "tsx", "c", "cpp", "h", "hpp", "java", "go", "rs", "rb", "php", "sql", "sh", "bat", "ps1", "css", "scss", "yaml", "yml", "json", "toml", "ini", "env"}
}

def get_file_category(filename):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    for category, extensions in ALLOWED_EXTENSIONS.items():
        if ext in extensions:
            return category
    return "other"

def encode_image_to_base64(filepath):
    """Encode an image file to a base64 data URI string."""
    mime_type, _ = mimetypes.guess_type(filepath)
    if not mime_type:
        mime_type = "image/jpeg"
    with open(filepath, "rb") as image_file:
        encoded = base64.b64encode(image_file.read()).decode("utf-8")
    return f"data:{mime_type};base64,{encoded}"

def extract_pdf_text(filepath, max_pages=50):
    """Extract text from a PDF file."""
    text_content = []
    try:
        reader = pypdf.PdfReader(filepath)
        total_pages = len(reader.pages)
        pages_to_read = min(total_pages, max_pages)
        
        for i in range(pages_to_read):
            page_text = reader.pages[i].extract_text()
            if page_text and page_text.strip():
                text_content.append(f"--- Page {i+1}/{total_pages} ---\n{page_text.strip()}")
        
        if total_pages > max_pages:
            text_content.append(f"\n[Note: PDF truncated. Showing first {max_pages} of {total_pages} pages.]")
            
        full_text = "\n\n".join(text_content)
        return full_text if full_text.strip() else "[Empty PDF or scanned image document without extractable text]"
    except Exception as e:
        return f"[Error extracting text from PDF: {str(e)}]"

def extract_docx_text(filepath):
    """Extract text from a Word DOCX file."""
    try:
        doc = docx.Document(filepath)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        tables_text = []
        for t_idx, table in enumerate(doc.tables):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(" | ".join(row_cells))
            if table_rows:
                tables_text.append(f"Table {t_idx+1}:\n" + "\n".join(table_rows))
                
        all_text = []
        if paragraphs:
            all_text.append("\n".join(paragraphs))
        if tables_text:
            all_text.append("\n\n" + "\n\n".join(tables_text))
            
        return "\n\n".join(all_text) if all_text else "[Empty Document]"
    except Exception as e:
        return f"[Error extracting text from Word document: {str(e)}]"

def extract_xlsx_text(filepath, max_rows=150):
    """Extract tabular text from Excel spreadsheet using openpyxl."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        sheets_text = []
        for sheetname in wb.sheetnames[:5]:  # Process up to first 5 sheets
            sheet = wb[sheetname]
            sheet_rows = []
            row_count = 0
            for row in sheet.iter_rows(values_only=True):
                if row_count >= max_rows:
                    sheet_rows.append(f"[... truncated after {max_rows} rows ...]")
                    break
                row_values = [str(val).strip() if val is not None else "" for val in row]
                if any(row_values):
                    sheet_rows.append(" | ".join(row_values))
                    row_count += 1
            if sheet_rows:
                sheets_text.append(f"### Sheet: {sheetname} ({row_count} rows)\n" + "\n".join(sheet_rows))
        return "\n\n".join(sheets_text) if sheets_text else "[Empty Spreadsheet]"
    except Exception as e:
        return f"[Error reading spreadsheet: {str(e)}]"

def extract_csv_text(filepath, max_rows=150):
    """Extract summary and sample rows from CSV/TSV."""
    try:
        with open(filepath, "r", encoding="utf-8", errors="replace") as f:
            sample_lines = []
            for i, line in enumerate(f):
                if i < max_rows:
                    sample_lines.append(line.rstrip())
                else:
                    break
        content = "\n".join(sample_lines)
        if i >= max_rows:
            content += f"\n[Showing first {max_rows} rows]"
        return content
    except Exception as e:
        return f"[Error extracting CSV: {str(e)}]"

def extract_text_file(filepath, max_lines=1500):
    """Extract plain text / code / config file contents."""
    try:
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                lines = f.readlines()
        except UnicodeDecodeError:
            with open(filepath, "r", encoding="latin-1", errors="replace") as f:
                lines = f.readlines()
                
        if len(lines) > max_lines:
            content = "".join(lines[:max_lines]) + f"\n\n[Truncated: showing first {max_lines} of {len(lines)} lines]"
        else:
            content = "".join(lines)
        return content
    except Exception as e:
        return f"[Error reading text file: {str(e)}]"

def process_uploaded_file(file_storage, groq_service=None):
    """
    Saves the uploaded Werkzeug FileStorage object, analyzes it,
    and returns a structured attachment dictionary.
    """
    raw_name = file_storage.filename or "uploaded_file"
    # Sanitize filename to prevent Windows path traversal / full path issues
    clean_name = os.path.basename(raw_name.replace("\\", "/")) or "uploaded_file"
    ext = clean_name.rsplit(".", 1)[-1].lower() if "." in clean_name else ""
    
    unique_name = f"{uuid.uuid4().hex}_{clean_name}"
    filepath = os.path.join(UPLOAD_FOLDER, unique_name)
    
    file_storage.save(filepath)
    file_size = os.path.getsize(filepath) if os.path.exists(filepath) else 0
    category = get_file_category(clean_name)
    mime_type, _ = mimetypes.guess_type(filepath)
    if not mime_type:
        mime_type = "application/octet-stream"
        
    attachment = {
        "id": uuid.uuid4().hex,
        "original_name": clean_name,
        "filename": unique_name,
        "file_path": filepath,
        "url": f"/uploads/{unique_name}",
        "file_type": category,
        "extension": ext,
        "size_bytes": file_size,
        "size_formatted": format_file_size(file_size),
        "mime_type": mime_type,
        "extracted_text": "",
        "base64_data": "",
        "audio_transcript": "",
        "metadata": {}
    }
    
    # Process by category
    try:
        if category == "image":
            attachment["base64_data"] = encode_image_to_base64(filepath)
            try:
                with Image.open(filepath) as img:
                    attachment["metadata"]["dimensions"] = f"{img.width}x{img.height}"
                    attachment["metadata"]["format"] = img.format
            except Exception:
                pass
                
        elif category == "document":
            if ext == "pdf":
                attachment["extracted_text"] = extract_pdf_text(filepath)
            elif ext in ("docx", "doc"):
                attachment["extracted_text"] = extract_docx_text(filepath)
            elif ext in ("xlsx", "xls"):
                attachment["extracted_text"] = extract_xlsx_text(filepath)
            elif ext in ("csv", "tsv"):
                attachment["extracted_text"] = extract_csv_text(filepath)
            else:
                attachment["extracted_text"] = extract_text_file(filepath)
                
        elif category == "code":
            attachment["extracted_text"] = extract_text_file(filepath)
            attachment["metadata"]["language"] = ext
            
        elif category == "audio":
            if groq_service and groq_service.is_configured():
                try:
                    transcript = groq_service.transcribe_audio(filepath)
                    attachment["audio_transcript"] = transcript
                    attachment["extracted_text"] = f"[Audio Transcription ({clean_name})]:\n{transcript}"
                except Exception as e:
                    attachment["audio_transcript"] = f"Transcription note: {str(e)}"
                    attachment["extracted_text"] = f"[Audio file: {clean_name}]"
            else:
                attachment["extracted_text"] = f"[Audio file uploaded: {clean_name}]"
                
        elif category == "video":
            if groq_service and groq_service.is_configured() and ext in ("mp4", "webm", "mov"):
                try:
                    transcript = groq_service.transcribe_audio(filepath)
                    attachment["audio_transcript"] = transcript
                    attachment["extracted_text"] = f"[Video Audio Transcription ({clean_name})]:\n{transcript}"
                except Exception:
                    attachment["extracted_text"] = f"[Video file uploaded: {clean_name}]"
            else:
                attachment["extracted_text"] = f"[Video file uploaded: {clean_name}]"
                
        else:
            # "other" category - attempt plain text extraction if readable
            try:
                text_attempt = extract_text_file(filepath, max_lines=200)
                if text_attempt and not text_attempt.startswith("[Error"):
                    attachment["extracted_text"] = text_attempt
                else:
                    attachment["extracted_text"] = f"[{clean_name}: Binary/Special format, {attachment['size_formatted']}]"
            except Exception:
                attachment["extracted_text"] = f"[{clean_name}: {attachment['size_formatted']}]"
                
    except Exception as e:
        attachment["extracted_text"] = f"[Error analyzing {clean_name}: {str(e)}]"
        
    return attachment

def format_file_size(bytes_size):
    """Format bytes to human readable string."""
    for unit in ['B', 'KB', 'MB', 'GB']:
        if bytes_size < 1024.0:
            return f"{bytes_size:.1f} {unit}"
        bytes_size /= 1024.0
    return f"{bytes_size:.1f} TB"
